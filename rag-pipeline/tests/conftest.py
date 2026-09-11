"""Fixture dùng chung cho mọi test: tạo tài liệu mẫu giống trang Fair Work."""
import io

import pytest

SAMPLE_HTML = """<!DOCTYPE html>
<html><head><title>Visa holders and migrants | Fair Work Ombudsman</title></head>
<body>
  <header><nav><a href="/">Home</a> <a href="/pay">Pay</a></nav></header>
  <div class="breadcrumb">Home > Find help for > Visa holders</div>
  <main>
    <h1>Visa holders and migrants</h1>
    <p>All workers in Australia have the same workplace rights, regardless of citizenship or visa status.
       Your employer must pay you at least the minimum wage for all hours you work.</p>
    <h2>Pay and entitlements</h2>
    <p>Visa holders are entitled to the minimum wage, penalty rates and pay slips.
       Your employer cannot pay you less because you are on a visa.</p>
    <ul>
      <li><p>You must receive a pay slip within one working day of being paid.</p></li>
      <li>Cash payments must still meet the minimum wage.</li>
    </ul>
    <h2>Your visa is protected</h2>
    <p>Your visa will not be cancelled if you contact the Fair Work Ombudsman for help.
       Only the Department of Home Affairs can cancel a visa.</p>
    <table><tr><th>Visa</th><th>Work limit</th></tr><tr><td>Student</td><td>48 hours per fortnight</td></tr></table>
    <div class="share-links">Share this page on Facebook</div>
    <p>Was this page helpful?</p>
  </main>
  <footer>Copyright Fair Work Ombudsman. Last updated 1 July 2026.</footer>
</body></html>"""


def make_pdf_bytes() -> bytes:
    """PDF 3 trang: tiêu đề cỡ chữ lớn, heading vừa, chữ thường, footer lặp lại."""
    import pymupdf

    doc = pymupdf.open()
    pages = [
        ("Work health and safety in Australia", "Your rights at work",
         "Every worker has the right to a safe workplace. Your employer must give you "
         "protective equipment before you start work, and you do not have to pay for it."),
        (None, "Getting hurt at work",
         "If you are injured at work you may be able to get workers compensation. "
         "Tell your employer and see a doctor as soon as possible."),
        (None, "Who can help",
         "Contact the work health and safety regulator in your state or territory. "
         "You can use a free interpreter by calling the Translating and Interpreting Service."),
    ]
    for i, (title, heading, body) in enumerate(pages, start=1):
        page = doc.new_page()
        y = 72
        if title:
            page.insert_text((72, y), title, fontsize=20)
            y += 40
        page.insert_text((72, y), heading, fontsize=14)
        y += 28
        page.insert_textbox(pymupdf.Rect(72, y, 520, y + 200), body, fontsize=10)
        page.insert_text((72, 800), "Safe Work Australia info sheet", fontsize=8)
        page.insert_text((500, 800), str(i), fontsize=8)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_heading("Sexual harassment: know your rights", level=0)  # level 0 = style Title
    document.add_paragraph(
        "Sexual harassment is unwelcome sexual behaviour that could make a person feel "
        "offended, humiliated or intimidated. It is against the law at work."
    )
    document.add_heading("What your employer must do", level=1)
    document.add_paragraph(
        "Employers have a positive duty to take reasonable steps to prevent sexual harassment."
    )
    document.add_paragraph("Make a complaint to the Commission", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Organisation", "Phone"
    table.cell(1, 0).text, table.cell(1, 1).text = "Australian Human Rights Commission", "1300 656 419"
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


@pytest.fixture
def sample_html() -> bytes:
    return SAMPLE_HTML.encode("utf-8")


@pytest.fixture
def sample_pdf() -> bytes:
    return make_pdf_bytes()


@pytest.fixture
def sample_docx() -> bytes:
    return make_docx_bytes()


PATH_SETTINGS = {
    "RAW_DIR": "raw", "EXTRACTED_DIR": "extracted", "CLEAN_DIR": "clean",
    "CHUNKS_FILE": "chunks.jsonl", "TAGGED_FILE": "tagged.jsonl",
    "TAG_REVIEW_FILE": "tag_review.csv", "EMBEDDED_FILE": "embedded.jsonl",
    "CACHE_DIR": "cache", "TAG_CACHE_FILE": "cache/tags.json",
    "EXTRACT_MANIFEST": "extracted/manifest.json", "CLEAN_MANIFEST": "clean/manifest.json",
    "DB_FILE": "rag.db", "REPORTS_DIR": "reports", "EVAL_FILE": "questions.jsonl",
}


@pytest.fixture
def tmp_data(tmp_path, monkeypatch):
    """Chuyển mọi đường dẫn dữ liệu vào thư mục tạm và bật chế độ Gemini giả."""
    from config import settings

    for name, rel in PATH_SETTINGS.items():
        monkeypatch.setattr(settings, name, tmp_path / rel)
    monkeypatch.setenv(settings.FAKE_GEMINI_ENV, "1")
    return tmp_path
